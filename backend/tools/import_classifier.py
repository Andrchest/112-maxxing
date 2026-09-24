#!/usr/bin/env python3
"""Classifier v_046_24 xlsx → `reference/classifier/v046_24.json` + `.columns.json`
(HLD 70 §70.6.1-§70.6.2, D18, A-1; REQ-3044, REQ-5701-REQ-5717).

Reads the organizer's workbook read-only with `openpyxl` (the `tools` dependency group — never a
`backend/app` import, A-13) and writes two byte-stable JSON files:

* `v046_24.columns.json` — the routing column map: every organisation column group of row 1 (the
  «Классификатор МЧС» group split into its row-2 services «Служба 101», «ОДС ПСЦ», «МГПСС»),
  keyed by an org id, with its `name_ru` and its sub-columns (source column letter, row-2/row-3
  label, and the `when` feature condition that label states), plus the feature-flag vocabulary;
* `v046_24.json` — the 24 categories (`group_no`, `group_ru`, and any routing note written on the
  category row) and one record per classifier row (§70.6.2): `code` (column E «Номер»),
  `group_no`/`group_ru` (the owning category), `features` (Признак 1-3, positional),
  `extra_features`, `final_type_ru`, `ekp35_ru`, `main_service` (column M, the source's own code,
  verbatim) and `routing` — per org, one `{when, value}` per sub-column, aligned with the columns
  file. An org whose cells are all empty in a row is left out of that row.

Cell values are the source text with outer whitespace stripped; the three spellings of the
notification marker («карточка-112», «Карточка-112», «карточка -112») are written as
`карточка-112`. Whether a cell *counts* as a notification is the reader's business (A-1: non-empty
and not «нет реагирования»), so «нет реагирования» is kept verbatim.

    uv run python backend/tools/import_classifier.py           # write both files + manifest
    uv run python backend/tools/import_classifier.py --check   # exit 1 when either is stale
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

if __package__ in (None, ""):  # run as a script: make `tools.*` importable
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tools.import_manifest import CLASSIFIER_SOURCE, REFERENCE_DIR, REPO_ROOT, write_manifest

CLASSIFIER_ID = "v046_24"
ROWS_PATH = REFERENCE_DIR / "classifier" / f"{CLASSIFIER_ID}.json"
COLUMNS_PATH = REFERENCE_DIR / "classifier" / f"{CLASSIFIER_ID}.columns.json"
SOURCE_PATH = REPO_ROOT / CLASSIFIER_SOURCE

HEADER_ROWS = 3
#: Record columns (REQ-3044 / REQ-5701-5703: v_046_24 has no «Сценарий реагирования», so «Главная
#: служба» is M and the routing matrix starts at N).
COL_CODE, COL_GROUP = 5, 6
COL_FEATURES = (7, 8, 9)
COL_EXTRA, COL_FINAL, COL_EKP35, COL_MAIN = 10, 11, 12, 13
FIRST_ROUTING_COL = 14

#: Row-1 (or, inside «Классификатор МЧС», row-2) organisation header → org id. Whitespace in the
#: header is collapsed before lookup. A header missing here fails the build.
ORG_IDS: dict[str, str] = {
    "Служба 101": "MCHS_SLUZHBA_101",
    "ОДС ПСЦ": "MCHS_ODS_PSC",
    "МГПСС": "MCHS_MGPSS",
    "Классификатор МВД": "MVD",
    "Классификатор СМП": "SMP",
    "Классификатор МОСГАЗ": "MOSGAZ",
    "ЦЭМП": "TSEMP",
    "Классификатор ФСБ": "FSB",
    "Классификатор Мособлгаз": "MOSOBLGAZ",
    "Автомобильные дороги": "AVTODOROGI",
    "Мосгортранс": "MOSGORTRANS",
    "Гор. Хозяйство": "GOR_KHOZYAYSTVO",
    "ГОРМОСТ": "GORMOST",
    "Канал имени Москвы": "KANAL_IMENI_MOSKVY",
    "МГТС": "MGTS",
    "Метро": "METRO",
    "Мосводоканал": "MOSVODOKANAL",
    "МОЭК": "MOEK",
    'МОЭСК (ПАО "Россети Московский регион")': "MOESK",
    "ОЭК": "OEK",
    "Мослифт": "MOSLIFT",
    "ЦОДД": "TSODD",
    "Деп. ЖКХ": "DEP_ZHKH",
    "Департамент РБиПК (ГКУ МОСБЕЗ)": "DEP_RBIPK",
    "Аппарат МЭРА": "APPARAT_MERA",
    "Москоллектор": "MOSKOLLEKTOR",
    "РЖД (МосковскаяЖД)": "RZHD",
    "Департамент образования": "DEP_EDUCATION",
    "Центррегионводхоз (Московско-Окское БВУ)": "TSENTRREGIONVODKHOZ",
    "Военная комендатура": "MILITARY_COMMANDANT",
    "ОАТИ": "OATI",
    "Мосводосток": "MOSVODOSTOK",
    "Департамент ППиООС": "DEP_PPIOOS",
    "ОД Департамент ТСЗН": "DEP_TSZN",
    "РСВО": "RSVO",
    "ЭВАЖД": "EVAZHD",
    "МСППН": "MSPPN",
    "ДТУ_Р (Ритуал)": "DTU_RITUAL",
    "ДТУ": "DTU",
    "Росгвардия": "ROSGVARDIYA",
    "Территориальные ОИВ": "TERRITORIAL_OIV",
    "Территориальные ОИВ ТиНАО": "TERRITORIAL_OIV_TINAO",
    "Автомобильные дороги АО г.Москвы": "AVTODOROGI_AO",
    "Департамент строительства города Москвы": "DEP_CONSTRUCTION",
    "Комитет ветеринарии": "COMMITTEE_VETERINARY",
    "Мосжилинспекция": "MOSZHILINSPEKTSIYA",
    "Департамент культуры": "DEP_CULTURE",
    "ГКУ ЦСА имени Е.П.Глинки": "GKU_TSSA",
    "ГКУ НТУ": "GKU_NTU",
    "ФСО": "FSO",
    "ГУП МСР": "GUP_MSR",
    "Комитет по туризму г.Москвы": "MOSTURIZM",
    "ДГП (Департамент градостроительной политики)": "DGP",
    "ЦУКБ Министерство обороны": "TSUKB_MO",
    "ЦУКБ.БПЛА Министерство обороны": "TSUKB_BPLA_MO",
    "ГКУ Организатор перевозок": "ORGANIZATOR_PEREVOZOK",
    "ГПБУ Мосэкомониторинг": "MOSEKOMONITORING",
    "Министерство обороны РХБЗ": "MO_RKHBZ",
    "ООО Ситиэнерго": "SITIENERGO",
    "Депортамент гражданского строительства": "DEP_CIVIL_CONSTRUCTION",
}

#: The feature flags the sub-column labels condition on, with the label wording they come from.
FLAGS: dict[str, str] = {
    "no_access": "НД - НЕТ ДОСТУПА",
    "threat_to_people": "УЛ - УГРОЗА ЛЮДЯМ / угроза людям",
    "casualties": "ПП - ПОСТРАДАВШИЕ ПОГИБШИЕ / Пострадавшие / пострадавшие/погибшие",
    "casualties_not_on_site": "Пострадавшие не на месте",
    "offence": "Правонарушение",
    "gasification": "газификация",
    "medical_help": "мед. помощь",
    "evacuation_required": "треб. Эвакуация",
    "over_five_people": ">5 чел / ОД",
    "traffic_blocked": "перекрытие движение",
    "tunnel": "тоннель",
    "pedestrian": "пеш",
    "road": "ав",
    "communication_object": "на объектах связи",
    "construction_site": "стройка",
    "listed_object": "объект из перечня",
}

_F, _T = False, True
#: (org id, sub-column label) → the `when` condition the label states. A label that names a
#: delivery channel rather than a feature («Дежурная служба АРМ-112», «(КУБ)», «Москва», …) and an
#: org with a single unlabelled column have `{}`: the cell applies unconditionally. A labelled
#: sub-column missing here fails the build.
WHEN: dict[tuple[str, str], dict[str, bool]] = {
    ("MCHS_SLUZHBA_101", "Служба 101 (признак НД - НЕТ ДОСТУПА не выбран)"): {"no_access": _F},
    ("MCHS_SLUZHBA_101", "Служба 101 (выбран признак НД - НЕТ ДОСТУПА)"): {"no_access": _T},
    ("MCHS_ODS_PSC", "ОДС ПСЦ (другие признаки не выбраны)"): {
        "threat_to_people": _F,
        "casualties": _F,
        "no_access": _F,
    },
    ("MCHS_ODS_PSC", "ОДС ПСЦ (выбран признак УЛ - УГРОЗА ЛЮДЯМ)"): {"threat_to_people": _T},
    ("MCHS_ODS_PSC", "ОДС ПСЦ (выбран признак ПП - ПОСТРАДАВШИЕ ПОГИБШИЕ)"): {"casualties": _T},
    ("MCHS_ODS_PSC", "ОДС ПСЦ (выбран признак НД - НЕТ ДОСТУПА)"): {"no_access": _T},
    ("MVD", "признак Правонарушение или Пострадавшие не выбран"): {"offence": _F, "casualties": _F},
    ("MVD", "выбран признак Правонарушение"): {"offence": _T},
    ("MVD", "выбран признак Пострадавшие"): {"casualties": _T},
    ("SMP", "Классификатор СМП (признак Пострадавшие не выбран)"): {"casualties": _F},
    ("SMP", "Классификатор СМП (выбран признак Пострадавшие)"): {"casualties": _T},
    ("SMP", "Классификатор СМП (выбран признак Пострадавшие не на месте)"): {
        "casualties_not_on_site": _T
    },
    ("MOSGAZ", "признак не выбран"): {"gasification": _F},
    ("MOSGAZ", "газификация"): {"gasification": _T},
    ("TSEMP", "признаки не выбраны"): {
        "threat_to_people": _F,
        "casualties": _F,
        "medical_help": _F,
        "evacuation_required": _F,
    },
    ("TSEMP", "угроза людям"): {"threat_to_people": _T},
    ("TSEMP", "пострадавшие/погибшие"): {"casualties": _T},
    ("TSEMP", "мед. помощь"): {"medical_help": _T},
    ("TSEMP", "треб. Эвакуация"): {"evacuation_required": _T},
    ("FSB", "признак не выбран"): {"over_five_people": _F},
    ("FSB", ">5 чел / ОД"): {"over_five_people": _T},
    ("MOSGORTRANS", "признак не выбран"): {"casualties": _F, "traffic_blocked": _F},
    ("MOSGORTRANS", "постр / погибшие"): {"casualties": _T},
    ("MOSGORTRANS", "перекрытие движение"): {"traffic_blocked": _T},
    ("GORMOST", "признак не выбран"): {"tunnel": _F, "pedestrian": _F, "road": _F},
    ("GORMOST", "тоннель"): {"tunnel": _T},
    ("GORMOST", "пеш"): {"pedestrian": _T},
    ("GORMOST", "ав"): {"road": _T},
    ("MGTS", "реагирование всегда"): {},
    ("MGTS", "на объектах связи"): {"communication_object": _T},
    ("DEP_RBIPK", "Дежурная служба АРМ-112"): {},
    ("DEP_RBIPK", "МКП, Аналитика (Старый КРИМ)"): {},
    ("DEP_CONSTRUCTION", "признак не выбран"): {"construction_site": _F},
    ("DEP_CONSTRUCTION", "стройка"): {"construction_site": _T},
    ("DEP_CULTURE", "объект из перечня"): {"listed_object": _T},
    ("GUP_MSR", "(КУБ)"): {},
    ("GUP_MSR", "пожары"): {},
    ("DGP", "интеграция"): {},
    ("DGP", "АРМ-112"): {},
    ("ORGANIZATOR_PEREVOZOK", "признак не выбран"): {"traffic_blocked": _F},
    ("ORGANIZATOR_PEREVOZOK", "перекрытие движение"): {"traffic_blocked": _T},
    ("MO_RKHBZ", "События по полигонам"): {},
    ("MO_RKHBZ", "Москва"): {},
}

NOTIFICATION_MARKER = "карточка-112"
_MARKER_SPELLINGS = frozenset({"карточка-112", "карточка -112"})


def _text(value: object) -> str | None:
    """A cell as stripped text; `None` for an empty or whitespace-only cell."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _label(value: object) -> str | None:
    """A header label with runs of whitespace (newlines included) collapsed to one space."""
    text = _text(value)
    return re.sub(r"\s+", " ", text) if text is not None else None


def _cell_value(value: object) -> str | None:
    text = _text(value)
    if text is not None and text.casefold() in _MARKER_SPELLINGS:
        return NOTIFICATION_MARKER
    return text


def _letter(column: int) -> str:
    from openpyxl.utils import get_column_letter

    return str(get_column_letter(column))


def _merged_spans(sheet: Any, row: int) -> dict[int, int]:
    """Row `row`'s merged ranges as first column → last column."""
    spans: dict[int, int] = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row == row:
            spans[merged.min_col] = merged.max_col
    return spans


def _orgs(sheet: Any) -> list[dict[str, Any]]:
    """The organisation column groups from N to the last column, in column order."""
    row1, row2 = _merged_spans(sheet, 1), _merged_spans(sheet, 2)
    orgs: list[dict[str, Any]] = []
    column = FIRST_ROUTING_COL
    while column <= sheet.max_column:
        last = row1.get(column, column)
        header = _label(sheet.cell(1, column).value)
        if header is None:
            raise ValueError(f"column {_letter(column)}: routing group without a row-1 header")
        # A group whose row 2 is itself merged into sub-groups (only «Классификатор МЧС») is split
        # into one org per row-2 header.
        if any(column <= start <= last and end > start for start, end in row2.items()):
            sub = column
            while sub <= last:
                sub_last = row2.get(sub, sub)
                name = _label(sheet.cell(2, sub).value)
                assert name is not None, f"{_letter(sub)}2: empty sub-group header"
                orgs.append(_org(sheet, name, sub, sub_last, label_row=3))
                sub = sub_last + 1
        else:
            label_row = 2 if any(sheet.cell(2, c).value for c in range(column, last + 1)) else 3
            orgs.append(_org(sheet, header, column, last, label_row=label_row))
        column = last + 1
    return orgs


def _org(sheet: Any, name: str, first: int, last: int, *, label_row: int) -> dict[str, Any]:
    org_id = ORG_IDS.get(name)
    if org_id is None:
        raise ValueError(f"{_letter(first)}: no org id for header «{name}»; add it to ORG_IDS")
    sub_columns = []
    for column in range(first, last + 1):
        label = _label(sheet.cell(label_row, column).value)
        if label is None:
            when: dict[str, bool] = {}
        elif (org_id, label) in WHEN:
            when = WHEN[(org_id, label)]
        else:
            raise ValueError(f"{_letter(column)}: no `when` for ({org_id}, «{label}»)")
        sub_columns.append({"column": _letter(column), "label_ru": label, "when": when})
    return {"org_id": org_id, "name_ru": name, "first": first, "sub_columns": sub_columns}


def _routing(row: tuple[Any, ...], orgs: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    routing: dict[str, list[dict[str, Any]]] = {}
    for org in orgs:
        cells = [
            {"when": sub["when"], "value": _cell_value(row[org["first"] - 1 + index])}
            for index, sub in enumerate(org["sub_columns"])
        ]
        if any(cell["value"] is not None for cell in cells):
            routing[org["org_id"]] = cells
    return routing


def build(source: Path = SOURCE_PATH) -> tuple[dict[str, Any], dict[str, Any]]:
    """`(rows document, columns document)` for the workbook at `source`."""
    from openpyxl import load_workbook

    # `data_only`: column E «Номер» is a formula (`=A*1000000+B*10000+C*100+D`); its cached
    # value is the code.
    workbook = load_workbook(source, data_only=True, read_only=False)
    sheet = workbook.worksheets[0]
    orgs = _orgs(sheet)

    groups: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=HEADER_ROWS + 1, values_only=True):
        if values[0] is None:  # a category-header row: no «Г» generation code
            group_ru = _text(values[COL_GROUP - 1])
            if group_ru is None:
                raise ValueError(f"category row {len(groups) + 1}: no category name")
            group_no = len(groups) + 1
            stated = values[COL_CODE - 1]
            if stated is not None and int(stated) != group_no:
                raise ValueError(f"category «{group_ru}» numbered {stated}, expected {group_no}")
            notes = {
                org["org_id"]: note
                for org in orgs
                for offset in range(len(org["sub_columns"]))
                if (note := _text(values[org["first"] - 1 + offset])) is not None
            }
            groups.append({"group_no": group_no, "group_ru": group_ru, "routing_notes": notes})
            continue
        if not groups:
            raise ValueError("a data row precedes the first category row")
        code = values[COL_CODE - 1]
        if not isinstance(code, int | float):
            raise ValueError(
                f"row after «{groups[-1]['group_ru']}»: «Номер» {code!r} is not cached"
            )
        # Признак 1-3 are positional: an empty middle one (six rows, e.g. «метро» / — / «дым»)
        # stays as "" so Признак 3 keeps its place; trailing empties are dropped.
        features = [_text(values[column - 1]) or "" for column in COL_FEATURES]
        while features and not features[-1]:
            features.pop()
        extra = _text(values[COL_EXTRA - 1])
        rows.append(
            {
                "code": str(int(code)),
                "group_no": groups[-1]["group_no"],
                "group_ru": groups[-1]["group_ru"],
                "features": features,
                "extra_features": [extra] if extra is not None else [],
                "final_type_ru": _text(values[COL_FINAL - 1]) or "",
                "ekp35_ru": _text(values[COL_EKP35 - 1]),
                "main_service": _text(values[COL_MAIN - 1]),
                "routing": _routing(values, orgs),
            }
        )
    codes = [row["code"] for row in rows]
    if len(set(codes)) != len(codes):
        raise ValueError("duplicate classifier codes")

    rows_document = {"classifier_id": CLASSIFIER_ID, "groups": groups, "rows": rows}
    columns_document = {
        "classifier_id": CLASSIFIER_ID,
        "source": {
            "path": CLASSIFIER_SOURCE,
            "sheet": sheet.title,
            "dimensions": sheet.dimensions,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "header_rows": HEADER_ROWS,
        },
        "record_columns": {
            "code": _letter(COL_CODE),
            "group": _letter(COL_GROUP),
            "features": [_letter(column) for column in COL_FEATURES],
            "extra_features": _letter(COL_EXTRA),
            "final_type_ru": _letter(COL_FINAL),
            "ekp35_ru": _letter(COL_EKP35),
            "main_service": _letter(COL_MAIN),
        },
        "notification_rule": (
            "a cell is a notification iff it is non-empty and not «нет реагирования» (A-1)"
        ),
        "flags": FLAGS,
        "orgs": {
            org["org_id"]: {"name_ru": org["name_ru"], "sub_columns": org["sub_columns"]}
            for org in orgs
        },
    }
    return rows_document, columns_document


def _dump_rows(document: dict[str, Any]) -> str:
    """One category or classifier row per line: ~2.6 MB instead of ~4 MB indented, and a
    regenerated row still shows up as one changed line in a diff."""
    lines = ["{", f'"classifier_id": {json.dumps(document["classifier_id"])},']
    for key in ("groups", "rows"):
        items = document[key]
        lines.append(f'"{key}": [')
        lines.extend(
            json.dumps(item, ensure_ascii=False, separators=(",", ":"))
            + ("," if index < len(items) - 1 else "")
            for index, item in enumerate(items)
        )
        lines.append("]," if key == "groups" else "]")
    lines.append("}")
    return "\n".join(lines) + "\n"


def render(source: Path = SOURCE_PATH) -> tuple[str, str]:
    """The byte-stable `(v046_24.json, v046_24.columns.json)` texts."""
    rows_document, columns_document = build(source)
    columns_text = json.dumps(columns_document, ensure_ascii=False, indent=1) + "\n"
    return _dump_rows(rows_document), columns_text


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(
        description="Build reference/classifier/*.json (HLD 70 §70.6.2)"
    )
    parser.add_argument("--check", action="store_true", help="exit 1 when either file is stale")
    args = parser.parse_args(argv)
    rows_text, columns_text = render()
    targets = ((ROWS_PATH, rows_text), (COLUMNS_PATH, columns_text))
    if args.check:
        stale = [
            path
            for path, text in targets
            if not path.is_file() or path.read_text(encoding="utf-8") != text
        ]
        for path in stale:
            print(f"{path} is stale; run backend/tools/import_classifier.py", file=sys.stderr)
        return 1 if stale else 0
    for path, text in targets:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
    write_manifest()
    print(f"wrote {ROWS_PATH.relative_to(REPO_ROOT)} and {COLUMNS_PATH.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
